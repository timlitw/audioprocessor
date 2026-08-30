"""Export a transcript project to document formats (txt, md, rtf, docx, srt)."""

from pathlib import Path

from core.project import TranscriptProject

FORMATS = ("docx", "rtf", "md", "txt", "srt")


def build_blocks(
    project: TranscriptProject, timestamps: bool, speakers: bool, types: bool
) -> list[tuple[str, str]]:
    """Turn segments into a flat list of (kind, text) blocks.

    kind is "header", "blank", or "line". Layout rules:
    - speaker header when the effective speaker changes (blank line before, except at top)
    - "Singing" header when singing starts; blank line when it ends
    - speech gets no header
    """
    blocks: list[tuple[str, str]] = []
    current_speaker = ""
    singing = False

    for i, seg in enumerate(project.segments):
        if not seg.text.strip():
            continue

        speaker = project.get_effective_speaker_label(i) if speakers else ""
        is_singing = types and seg.type == "singing"

        speaker_changed = speakers and speaker and speaker != current_speaker
        singing_started = is_singing and not singing
        singing_ended = singing and not is_singing

        if (speaker_changed or singing_started or singing_ended) and blocks:
            blocks.append(("blank", ""))
        if speaker_changed:
            blocks.append(("header", speaker))
            current_speaker = speaker
        if singing_started:
            blocks.append(("header", "Singing"))
        singing = is_singing

        text = seg.text.strip()
        if timestamps:
            text = f"[{project.format_time(seg.start)}] {text}"
        blocks.append(("line", text))

    return blocks


def _title(project: TranscriptProject) -> str:
    return Path(project.audio_file).stem if project.audio_file else "Transcript"


# --- Renderers ---

def _render_txt(blocks) -> str:
    return "\n".join(text for _, text in blocks) + "\n"


def _render_md(project, blocks) -> str:
    out = [f"# {_title(project)}", ""]
    for kind, text in blocks:
        if kind == "header":
            out.append(f"**{text}**  ")
        elif kind == "blank":
            out.append("")
        else:
            out.append(f"{text}  ")
    return "\n".join(out) + "\n"


def _rtf_escape(text: str) -> str:
    out = []
    for ch in text:
        if ch in "\\{}":
            out.append("\\" + ch)
        elif ord(ch) < 128:
            out.append(ch)
        else:
            code = ord(ch)
            if code > 32767:
                code -= 65536
            out.append(f"\\u{code}?")
    return "".join(out)


def _render_rtf(project, blocks) -> str:
    out = [r"{\rtf1\ansi\deff0{\fonttbl{\f0 Calibri;}}\fs22"]
    out.append(r"{\b\fs28 " + _rtf_escape(_title(project)) + r"}\par\par")
    for kind, text in blocks:
        if kind == "header":
            out.append(r"{\b " + _rtf_escape(text) + r"}\par")
        elif kind == "blank":
            out.append(r"\par")
        else:
            out.append(_rtf_escape(text) + r"\par")
    out.append("}")
    return "\n".join(out)


def _write_docx(project, blocks, path: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(_title(project), level=1)
    for kind, text in blocks:
        if kind == "header":
            doc.add_paragraph().add_run(text).bold = True
        elif kind == "blank":
            doc.add_paragraph()
        else:
            doc.add_paragraph(text)
    doc.save(path)


def _srt_time(seconds: float) -> str:
    ms = int(round(seconds * 1000))
    h, rem = divmod(ms, 3_600_000)
    m, rem = divmod(rem, 60_000)
    s, ms = divmod(rem, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def _render_srt(project, speakers: bool) -> tuple[str, int]:
    out = []
    n = 0
    current_speaker = ""
    for i, seg in enumerate(project.segments):
        text = seg.text.strip()
        if not text:
            continue
        if speakers:
            speaker = project.get_effective_speaker_label(i)
            if speaker and speaker != current_speaker:
                text = f"{speaker}: {text}"
                current_speaker = speaker
        n += 1
        out.append(f"{n}\n{_srt_time(seg.start)} --> {_srt_time(seg.end)}\n{text}\n")
    return "\n".join(out), n


def export_transcript(
    project: TranscriptProject,
    path: str,
    fmt: str,
    timestamps: bool,
    speakers: bool,
    types: bool,
) -> int:
    """Write the transcript to `path` in `fmt`. Returns number of text lines written."""
    if fmt == "srt":
        content, n = _render_srt(project, speakers)
        Path(path).write_text(content, encoding="utf-8")
        return n

    blocks = build_blocks(project, timestamps, speakers, types)
    n = sum(1 for kind, _ in blocks if kind == "line")

    if fmt == "txt":
        Path(path).write_text(_render_txt(blocks), encoding="utf-8")
    elif fmt == "md":
        Path(path).write_text(_render_md(project, blocks), encoding="utf-8")
    elif fmt == "rtf":
        Path(path).write_text(_render_rtf(project, blocks), encoding="ascii")
    elif fmt == "docx":
        _write_docx(project, blocks, path)
    else:
        raise ValueError(f"Unknown export format: {fmt}")
    return n
